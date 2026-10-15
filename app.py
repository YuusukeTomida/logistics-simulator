import streamlit as st
import pandas as pd
import numpy as np
import cv2
import os
import math
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenPyxlImage

# 1. ページ基本設定
st.set_page_config(
    page_title="配達エリア・トンキロ最適化シミュレーター（大阪府版）",
    page_icon="🚚",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.title("🚚 運送2社 営業所受持エリア・トンキロ最適化シミュレーター（大阪府版）")
st.caption("大阪府の市区町村別・会社別の受持選択を行い、「現状」「4つの改正案」「最適解」のトンキロ・配達重量の変化を一括比較・シミュレーションします。")

# 3. 役所・役場座標 ＆ 営業所初期座標定義
MUNICIPAL_HALL_COORDS = {
    '大阪市都島区': (34.7011, 135.5303), '大阪市福島区': (34.6922, 135.4727), '大阪市此花区': (34.6853, 135.4619),
    '大阪市西区': (34.6756, 135.4872), '大阪市港区': (34.6644, 135.4608), '大阪市大正区': (34.6528, 135.4711),
    '大阪市天王寺区': (34.6561, 135.5208), '大阪市浪速区': (34.6558, 135.4981), '大阪市西淀川区': (34.7119, 135.4578),
    '大阪市東淀川区': (34.7394, 135.5347), '大阪市東成区': (34.6694, 135.5472), '大阪市生野区': (34.6550, 135.5439),
    '大阪市旭区': (34.7214, 135.5428), '大阪市城東区': (34.6989, 135.5458), '大阪市阿倍野区': (34.6386, 135.5175),
    '大阪市住吉区': (34.6086, 135.4958), '大阪市東住吉区': (34.6228, 135.5275), '大阪市西成区': (34.6358, 135.4947),
    '大阪市淀川区': (34.7203, 135.4831), '大阪市鶴見区': (34.7036, 135.5714), '大阪市住之江区': (34.6097, 135.4722),
    '大阪市平野区': (34.6214, 135.5489), '大阪市北区': (34.7056, 135.5103), '大阪市中央区': (34.6822, 135.5108),
    '堺市堺区': (34.5731, 135.4831), '堺市中区': (34.5369, 135.4967), '堺市東区': (34.5328, 135.5222),
    '堺市西区': (34.5483, 135.4608), '堺市南区': (34.4925, 135.5019), '堺市北区': (34.5661, 135.5133),
    '堺市美原区': (34.5269, 135.5583), '岸和田市': (34.4597, 135.3719), '豊中市': (34.7814, 135.4703),
    '池田市': (34.8214, 135.4278), '吹田市': (34.7583, 135.5167), '泉大津市': (34.5028, 135.4056),
    '高槻市': (34.8483, 135.6181), '貝塚市': (34.4369, 135.3583), '守口市': (34.7358, 135.5653),
    '枚方市': (34.8161, 135.6508), '茨木市': (34.8161, 135.5683), '八尾市': (34.6269, 135.6008),
    '泉佐野市': (34.4086, 135.3283), '富田林市': (34.5008, 135.6008), '寝屋川市': (34.7656, 135.6269),
    '河内長野市': (34.4558, 135.5639), '松原市': (34.5778, 135.5539), '大東市': (34.7119, 135.6239),
    '和泉市': (34.4883, 135.4267), '箕面市': (34.8269, 135.4703), '柏原市': (34.5808, 135.6289),
    '羽曳野市': (34.5583, 135.6067), '門真市': (34.7336, 135.5886), '摂津市': (34.7792, 135.5622),
    '高石市': (34.5208, 135.4383), '藤井寺市': (34.5739, 135.5967), '東大阪市': (34.6794, 135.6008),
    '泉南市': (34.3639, 135.2858), '四條畷市': (34.7369, 135.6419), '交野市': (34.7878, 135.6886),
    '大阪狭山市': (34.5036, 135.5539), '阪南市': (34.3583, 135.2389), '島本町': (34.8869, 135.6639),
    '豊能町': (34.9383, 135.4739), '能勢町': (34.9669, 135.3989), '忠岡町': (34.4869, 135.3919),
    '熊取町': (34.4036, 135.3539), '田尻町': (34.3983, 135.2919), '岬町': (34.3214, 135.1539),
    '太子町': (34.5222, 135.6489), '河南町': (34.4883, 135.6358), '千早赤阪村': (34.4569, 135.6189)
}

MUNICIPAL_ADDRESS_LOOKUP = {
    '大阪市都島区': '大阪府大阪市都島区中野町5-15-21',
    '堺市堺区': '大阪府堺市堺区南瓦町3-1',
    '岸和田市': '大阪府岸和田市岸城町7-1',
    '東大阪市': '大阪府東大阪市荒本北1-1-1',
    '堺市美原区': '大阪府堺市美原区黒山167-1',
    '茨木市': '大阪府茨木市駅前3-8-13',
    '泉佐野市': '大阪府泉佐野市市場東1-1-1',
    '松原市': '大阪府松原市阿保1-1-1',
    '羽曳野市': '大阪府羽曳野市誉田4-1-1'
}

EXISTING_OFFICES_INFO = {
    ('J社', '大阪'): {'address': '大阪府茨木市宿久庄2-10-2', 'coords': (34.8583, 135.5383)},
    ('J社', '尼崎'): {'address': '兵庫県尼崎市西高洲町16-19', 'coords': (34.7083, 135.4083)},
    ('J社', '大阪南港'): {'address': '大阪府大阪市住之江区南港南1-1-125', 'coords': (34.6183, 135.4083)},
    ('J社', '東大阪'): {'address': '大阪府八尾市西高安町4-57-1', 'coords': (34.6308, 135.6383)},
    ('J社', '松原'): {'address': '大阪府松原市三宅西6-891-1', 'coords': (34.5883, 135.5383)},
    ('J社', '貝塚'): {'address': '大阪府貝塚市港17-8', 'coords': (34.4483, 135.3483)},
    ('J社', '京都'): {'address': '京都府久世郡久御山町下津屋上ノ浜28-2', 'coords': (34.8883, 135.7383)},
    ('T社', '大阪中央'): {'address': '大阪府大阪市鶴見区焼野3-2-11', 'coords': (34.7139, 135.5839)},
    ('T社', '尼崎'): {'address': '兵庫県尼崎市東海岸町21-10', 'coords': (34.6983, 135.4183)},
    ('T社', '南大阪'): {'address': '大阪府堺市堺区築港八幡町1-1', 'coords': (34.6083, 135.4383)},
    ('T社', '北大阪'): {'address': '大阪府茨木市宮島2-5-1', 'coords': (34.8039, 135.5839)},
    ('T社', '東大阪'): {'address': '大阪府東大阪市本庄中1-4-90', 'coords': (34.6883, 135.6083)},
    ('T社', '泉佐野'): {'address': '大阪府泉佐野市下瓦屋町2-2-50', 'coords': (34.4239, 135.3339)},
    ('T社', '京都'): {'address': '京都府京都市伏見区横大路一本木22', 'coords': (34.9283, 135.7483)},
}

FULL_CITY_COORDS_LOOKUP = {
    '東大阪市': (34.6794, 135.6008), '東大阪': (34.6794, 135.6008),
    '大阪市都島区': (34.7011, 135.5303), '都島区': (34.7011, 135.5303), '都島': (34.7011, 135.5303),
    '大阪市福島区': (34.6922, 135.4727), '福島区': (34.6922, 135.4727),
    '大阪市此花区': (34.6853, 135.4619), '此花区': (34.6853, 135.4619),
    '大阪市西区': (34.6756, 135.4872), '大阪市港区': (34.6644, 135.4608),
    '大阪市大正区': (34.6528, 135.4711), '大正区': (34.6528, 135.4711),
    '大阪市天王寺区': (34.6561, 135.5208), '天王寺区': (34.6561, 135.5208), '天王寺': (34.6561, 135.5208),
    '大阪市浪速区': (34.6558, 135.4981), '浪速区': (34.6558, 135.4981),
    '大阪市西淀川区': (34.7119, 135.4578), '西淀川区': (34.7119, 135.4578),
    '大阪市東淀川区': (34.7394, 135.5347), '東淀川区': (34.7394, 135.5347),
    '大阪市東成区': (34.6694, 135.5472), '東成区': (34.6694, 135.5472),
    '大阪市生野区': (34.6550, 135.5439), '生野区': (34.6550, 135.5439),
    '大阪市旭区': (34.7214, 135.5428), '大阪市城東区': (34.6989, 135.5458),
    '城東区': (34.6989, 135.5458), '大阪市阿倍野区': (34.6386, 135.5175),
    '阿倍野区': (34.6386, 135.5175), '阿倍野': (34.6386, 135.5175),
    '大阪市住吉区': (34.6086, 135.4958), '大阪市東住吉区': (34.6228, 135.5275),
    '東住吉区': (34.6228, 135.5275), '大阪市西成区': (34.6358, 135.4947),
    '西成区': (34.6358, 135.4947), '大阪市淀川区': (34.7203, 135.4831),
    '淀川区': (34.7203, 135.4831), '大阪市鶴見区': (34.7036, 135.5714),
    '鶴見区': (34.7036, 135.5714), '大阪市住之江区': (34.6097, 135.4722),
    '住之江区': (34.6097, 135.4722), '住之江': (34.6097, 135.4722),
    '大阪市平野区': (34.6214, 135.5489), '平野区': (34.6214, 135.5489),
    '大阪市北区': (34.7056, 135.5103), '大阪市中央区': (34.6822, 135.5108),
    '堺市堺区': (34.5731, 135.4831), '堺区': (34.5731, 135.4831),
    '堺市中区': (34.5369, 135.4967), '堺市東区': (34.5328, 135.5222),
    '堺市西区': (34.5483, 135.4608), '堺市南区': (34.4925, 135.5019),
    '堺市北区': (34.5661, 135.5133), '堺市美原区': (34.5269, 135.5583),
    '美原区': (34.5269, 135.5583), '美原': (34.5269, 135.5583),
    '岸和田市': (34.4597, 135.3719), '岸和田': (34.4597, 135.3719),
    '豊中市': (34.7814, 135.4703), '豊中': (34.7814, 135.4703),
    '池田市': (34.8214, 135.4278), '池田': (34.8214, 135.4278),
    '吹田市': (34.7583, 135.5167), '吹田': (34.7583, 135.5167),
    '泉大津市': (34.5028, 135.4056), '泉大津': (34.5028, 135.4056),
    '高槻市': (34.8483, 135.6181), '高槻': (34.8483, 135.6181),
    '貝塚市': (34.4369, 135.3583), '貝塚': (34.4369, 135.3583),
    '守口市': (34.7358, 135.5653), '守口': (34.7358, 135.5653),
    '枚方市': (34.8161, 135.6508), '枚方': (34.8161, 135.6508),
    '茨木市': (34.8161, 135.5683), '茨木': (34.8161, 135.5683),
    '八尾市': (34.6269, 135.6008), '八尾': (34.6269, 135.6008),
    '泉佐野市': (34.4086, 135.3283), '泉佐野': (34.4086, 135.3283),
    '富田林市': (34.5008, 135.6008), '富田林': (34.5008, 135.6008),
    '寝屋川市': (34.7656, 135.6269), '寝屋川': (34.7656, 135.6269),
    '河内長野市': (34.4558, 135.5639), '河内長野': (34.4558, 135.5639),
    '松原市': (34.5778, 135.5539), '松原': (34.5778, 135.5539),
    '大東市': (34.7119, 135.6239), '大東': (34.7119, 135.6239),
    '和泉市': (34.4883, 135.4267), '和泉': (34.4883, 135.4267),
    '箕面市': (34.8269, 135.4703), '箕面': (34.8269, 135.4703),
    '柏原市': (34.5808, 135.6289), '柏原': (34.5808, 135.6289),
    '羽曳野市': (34.5583, 135.6067), '羽曳野': (34.5583, 135.6067),
    '門真市': (34.7336, 135.5886), '門真': (34.7336, 135.5886),
    '摂津市': (34.7792, 135.5622), '摂津': (34.7792, 135.5622),
    '高石市': (34.5208, 135.4383), '高石': (34.5208, 135.4383),
    '藤井寺市': (34.5739, 135.5967), '藤井寺': (34.5739, 135.5967),
    '泉南市': (34.3639, 135.2858), '泉南': (34.3639, 135.2858),
    '四條畷市': (34.7369, 135.6419), '四條畷': (34.7369, 135.6419),
    '交野市': (34.7878, 135.6886), '交野': (34.7878, 135.6886),
    '大阪狭山市': (34.5036, 135.5539), '大阪狭山': (34.5036, 135.5539),
    '阪南市': (34.3583, 135.2389), '阪南': (34.3583, 135.2389),
    '島本町': (34.8869, 135.6639), '島本': (34.8869, 135.6639),
    '豊能町': (34.9383, 135.4739), '豊能': (34.9383, 135.4739),
    '能勢町': (34.9669, 135.3989), '能勢': (34.9669, 135.3989),
    '忠岡町': (34.4869, 135.3919), '忠岡': (34.4869, 135.3919),
    '熊取町': (34.4036, 135.3539), '熊取': (34.4036, 135.3539),
    '田尻町': (34.3983, 135.2919), '田尻': (34.3983, 135.2919),
    '岬町': (34.3214, 135.1539), '太子町': (34.5222, 135.6489),
    '河南町': (34.4883, 135.6358), '千早赤阪村': (34.4569, 135.6189),
    '千早赤阪': (34.4569, 135.6189), '尼崎市': (34.7083, 135.4083),
    '尼崎': (34.7083, 135.4083), '神戸市': (34.6901, 135.1955),
    '神戸': (34.6901, 135.1955), '西宮市': (34.7376, 135.3414),
    '西宮': (34.7376, 135.3414), '京都市': (34.9858, 135.7583),
    '京都': (34.9858, 135.7583), '堺市': (34.5731, 135.4831),
    '堺': (34.5731, 135.4831), '大阪市': (34.6853, 135.5208),
    '大阪': (34.6853, 135.5208)
}

def geocode_address_precise(address_str):
    if not address_str:
        return (34.6853, 135.5208)
    for kw, coords in sorted(FULL_CITY_COORDS_LOOKUP.items(), key=lambda x: len(x[0]), reverse=True):
        if kw in address_str:
            return coords
    return (34.6853, 135.5208)

def calc_haversine_distance(lat1, lon1, lat2, lon2):
    R = 6371.0
    dlat = math.radians(lat2 - lat1)
    dlon = math.radians(lon2 - lon1)
    a = math.sin(dlat / 2)**2 + math.cos(math.radians(lat1)) * math.cos(math.radians(lat2)) * math.sin(dlon / 2)**2
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return round(R * c * 1.2, 1)

def parse_pattern_num(val):
    s = str(val).strip()
    res = []
    for p in [1, 2, 3, 4]:
        if str(p) in s or f"案{p}" in s or f"案{chr(ord('１')+p-1)}" in s:
            res.append(p)
    if not res:
        if '全' in s or '一括' in s or 'ALL' in s.upper():
            return [1, 2, 3, 4]
    return res

# 2. 基本データファイルアップロード
st.sidebar.header("📁 基本データ読み込み")
uploaded_file = st.sidebar.file_uploader("「営業所受持地域.xlsx」をドラッグ＆ドロップ", type=["xlsx"], key="base_excel_uploader")

if uploaded_file is None:
    st.info("👈 左側のサイドバーから「営業所受持地域.xlsx」をアップロードしてシミュレーションを開始してください。")
    st.stop()

@st.cache_data
def load_data(file):
    xls = pd.ExcelFile(file)
    df_j = pd.read_excel(xls, sheet_name='J')
    df_t = pd.read_excel(xls, sheet_name='T')
    return df_j, df_t

try:
    df_j, df_t = load_data(uploaded_file)
except Exception as e:
    st.error("データの読み込みに失敗しました。J, T シートが含まれるExcelファイルであることを確認してください。")
    st.stop()

# 取り込み前のファイル事前ダウンロードボタンの設置
st.sidebar.markdown("---")
st.sidebar.header("📄 入力用Excelファイルの準備")
st.sidebar.caption("設定変更や一括取り込みを行う前に、テンプレートや基本データを事前にダウンロードできます。")

def create_full_import_template(df_j_in, df_t_in):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as t_writer:
        df_cust_sample = pd.DataFrame([
            {'対象パターン': '案1', '所在地住所': '大阪府茨木市駅前1-1-1'},
            {'対象パターン': '案2', '所在地住所': '大阪府堺市堺区南瓦町3-1'},
            {'対象パターン': '案3', '所在地住所': '大阪府東大阪市荒本北1-1-1'},
            {'対象パターン': '案4', '所在地住所': '大阪府岸和田市岸城町7-1'},
        ])
        df_cust_sample.to_excel(t_writer, sheet_name='新規営業所一覧', index=False)
        
        tmpl_rows = []
        for i in range(len(df_j_in)):
            code = df_j_in.loc[i, '市区町村コード']
            city = df_j_in.loc[i, '市区町村名\n（漢字）']
            off_j = df_j_in.loc[i, 'J社営業所名']
            off_t = df_t_in.loc[i, 'T社営業所名']
            
            tmpl_rows.append({
                '市町村コード': code,
                '市町村名': city,
                '市町村一括担当': 'なし',
                '市町村一括担当営業所': '-',
                'J社個別': off_j,
                'T社個別': off_t
            })
        df_tmpl_base = pd.DataFrame(tmpl_rows)
        for p_i in [1, 2, 3, 4]:
            df_tmpl_base.to_excel(t_writer, sheet_name=f'案{p_i}', index=False)
            
    return buf.getvalue()

st.sidebar.download_button(
    label="📄 一括取り込み用Excelテンプレートをダウンロード（現行営業所参照・案1〜4全シート付き）",
    data=create_full_import_template(df_j, df_t),
    file_name="変更案設定_一括取り込み用テンプレート.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True
)

st.sidebar.markdown("---")
st.sidebar.header("📥 変更案設定Excelを取り込み")

# 新規営業所の登録状態をパターン別に完全分離管理（custom_offices_p1 〜 custom_offices_p4）
for p_idx in [1, 2, 3, 4]:
    key_name = f"custom_offices_p{p_idx}"
    if key_name not in st.session_state:
        st.session_state[key_name] = []

def generate_office_colors(df_j, df_t, custom_offices):
    j_offices = sorted(list(set(df_j['J社営業所名'].dropna().astype(str).tolist() + [o['name'] for o in custom_offices if o['company'] == 'J社'])))
    t_offices = sorted(list(set(df_t['T社営業所名'].dropna().astype(str).tolist() + [o['name'] for o in custom_offices if o['company'] == 'T社'])))
    
    j_color_map = {}
    t_color_map = {}
    
    n_j = len(j_offices)
    for idx, off in enumerate(j_offices):
        f = 0.45 + 0.55 * (idx / max(n_j - 1, 1)) if n_j > 1 else 0.85
        j_color_map[off] = (int(255 * f), int(20 * (1.0 - f)), int(20 * (1.0 - f)))
        
    n_t = len(t_offices)
    for idx, off in enumerate(t_offices):
        f = 0.45 + 0.55 * (idx / max(n_t - 1, 1)) if n_t > 1 else 0.85
        t_color_map[off] = (int(30 * (1.0 - f)), int(140 * f), int(255 * f))
        
    return j_color_map, t_color_map, j_offices, t_offices

# ピクセル座標定義（完全正数化）
CITY_SEEDS = {
    '能勢町': (542, 143), '豊能町': (635, 232), '島本町': (897, 293), '高槻市': (816, 331), '枚方市': (923, 331),
    '池田市': (553, 351), '箕面市': (618, 359), '茨木市': (702, 292), '交野市': (932, 517), '豊中市': (592, 473),
    '吹田市': (702, 497), '摂津市': (750, 517), '寝屋川市': (832, 504), '守口市': (750, 583), '門真市': (818, 584),
    '四條畷市': (923, 584), '大東市': (818, 600), '東大阪市': (867, 672), '大阪市西淀川区': (566, 619),
    '大阪市東淀川区': (670, 550), '大阪市淀川区': (603, 593), '大阪市北区': (637, 625), '大阪市都島区': (680, 594),
    '大阪市旭区': (715, 588), '大阪市城東区': (714, 638), '大阪市鶴見区': (761, 619), '大阪市福島区': (594, 647),
    '大阪市此花区': (545, 673), '大阪市西区': (617, 673), '大阪市中央区': (653, 671), '大阪市東成区': (719, 676),
    '大阪市生野区': (708, 712), '大阪市港区': (555, 713), '大阪市浪速区': (628, 702), '大阪市天王寺区': (665, 707),
    '大阪市大正区': (582, 740), '大阪市西成区': (620, 736), '大阪市阿倍野区': (660, 741), '大阪市平野区': (724, 795),
    '大阪市住之江区': (533, 748), '大阪市住吉区': (660, 799), '大阪市東住吉区': (680, 799), '八尾市': (818, 769),
    '柏原市': (881, 837), '松原市': (723, 842), '藤井寺市': (817, 853), '羽曳野市': (811, 911), '太子町': (890, 962),
    '富田林市': (811, 1011), '大阪狭山市': (724, 1011), '河南町': (866, 1011), '千早赤阪村': (881, 1104),
    '河内長野市': (743, 1147), '堺市堺区': (607, 848), '堺市北区': (665, 850), '堺市西区': (565, 930),
    '堺市中区': (640, 930), '堺市東区': (682, 911), '堺市美原区': (754, 932), '堺市南区': (641, 1011),
    '高石市': (535, 930), '泉大津市': (504, 1000), '忠岡町': (504, 1011), '和泉市': (565, 1096),
    '岸和田市': (504, 1179), '貝塚市': (494, 1220), '熊取町': (426, 1179), '泉佐野市': (370, 1179),
    '田尻町': (308, 1187), '泉南市': (318, 1251), '阪南市': (239, 1280), '岬町': (118, 1336)
}

EXTRA_SEEDS = {
    '大阪市此花区': [(477, 710), (503, 713)],
    '泉佐野市': [(236, 1101)]
}

# 4. 距離・トンキロ算出関数
def compute_distance_and_tonkm(city_name, bulk_comp, bulk_off, j_indiv_off, t_indiv_off, wt_j_kg, wt_t_kg, office_info, municipal_coords):
    c_lat, c_lon = municipal_coords.get(city_name, (34.6853, 135.5208))
    
    if bulk_comp in ['なし', '-', 'nan', 'None', '']:
        bulk_off = '-'

    if bulk_off != '-' and bulk_off is not None:
        c_prefix = bulk_comp if bulk_comp in ['J社', 'T社'] else 'J社'
        coords = office_info.get((c_prefix, bulk_off), {}).get('coords')
        if not coords:
            other_prefix = 'T社' if c_prefix == 'J社' else 'J社'
            coords = office_info.get((other_prefix, bulk_off), {}).get('coords')
            
        if coords:
            dist = calc_haversine_distance(coords[0], coords[1], c_lat, c_lon)
            dist_j = dist
            dist_t = dist
        else:
            dist_j = 0.0
            dist_t = 0.0
    else:
        if j_indiv_off != '-' and j_indiv_off is not None:
            j_coords = office_info.get(('J社', j_indiv_off), {}).get('coords')
            if not j_coords:
                j_coords = office_info.get(('T社', j_indiv_off), {}).get('coords')
            dist_j = calc_haversine_distance(j_coords[0], j_coords[1], c_lat, c_lon) if j_coords else 0.0
        else:
            dist_j = 0.0
            
        if t_indiv_off != '-' and t_indiv_off is not None:
            t_coords = office_info.get(('T社', t_indiv_off), {}).get('coords')
            if not t_coords:
                t_coords = office_info.get(('J社', t_indiv_off), {}).get('coords')
            dist_t = calc_haversine_distance(t_coords[0], t_coords[1], c_lat, c_lon) if t_coords else 0.0
        else:
            dist_t = 0.0
            
    tk_j = round(dist_j * (wt_j_kg / 1000.0), 2)
    tk_t = round(dist_t * (wt_t_kg / 1000.0), 2)
    
    return dist_j, tk_j, dist_t, tk_t

def build_initial_scenario_df(office_info_dict):
    rows = []
    for i in range(len(df_j)):
        code = df_j.loc[i, '市区町村コード']
        city = df_j.loc[i, '市区町村名\n（漢字）']
        wt_j_kg = df_j.loc[i, '配達重量\n(日当たり)']
        wt_t_kg = df_t.loc[i, '配達重量\n(日当たり)']
        off_j_def, off_t_def = df_j.loc[i, 'J社営業所名'], df_t.loc[i, 'T社営業所名']
        
        dist_j, tk_j, dist_t, tk_t = compute_distance_and_tonkm(city, 'なし', '-', off_j_def, off_t_def, wt_j_kg, wt_t_kg, office_info_dict, MUNICIPAL_HALL_COORDS)
        
        rows.append({
            '市町村コード': code, '市町村名': city,
            '市町村一括担当': 'なし', '市町村一括担当営業所': '-',
            'J社個別': off_j_def, 'T社個別': off_t_def,
            'J社 距離(km)': dist_j, 'J社 配達重量(kg)': wt_j_kg, 'J社 トンキロ(t・km)': tk_j,
            'T社 距離(km)': dist_t, 'T社 配達重量(kg)': wt_t_kg, 'T社 トンキロ(t・km)': tk_t,
            '_j_def_off': off_j_def, '_t_def_off': off_t_def
        })
    return pd.DataFrame(rows)

for p_idx in [1, 2, 3, 4]:
    key_name = f"table_data_p{p_idx}"
    if key_name not in st.session_state:
        st.session_state[key_name] = build_initial_scenario_df(EXISTING_OFFICES_INFO).copy(deep=True)

# 条件（T社廃止拠点＝「東大阪」固定、現状が「委託」エリアは変更不可）を組み込んだ全体最適化エンジン
@st.cache_data
def solve_optimal_scenario_cached(df_j_in, df_t_in):
    j_existing_list = [k[1] for k in EXISTING_OFFICES_INFO.keys() if k[0] == 'J社']
    t_close_fixed = '東大阪'
    
    best_total_tk = float('inf')
    best_config = None
    
    cities = []
    for i in range(len(df_j_in)):
        c_name = df_j_in.loc[i, '市区町村名\n（漢字）']
        c_code = df_j_in.loc[i, '市区町村コード']
        wt_j = df_j_in.loc[i, '配達重量\n(日当たり)']
        wt_t = df_t_in.loc[i, '配達重量\n(日当たり)']
        off_j_def = str(df_j_in.loc[i, 'J社営業所名']).strip()
        off_t_def = str(df_t_in.loc[i, 'T社営業所名']).strip()
        
        is_j_sub = ('委託' in off_j_def) or (str(df_j_in.loc[i, 'J社外部委託地域']).strip() == '◯' if 'J社外部委託地域' in df_j_in.columns else False)
        is_t_sub = ('委託' in off_t_def) or (str(df_t_in.loc[i, 'T社外部委託地域']).strip() == '◯' if 'T社外部委託地域' in df_t_in.columns else False)
        is_sub = is_j_sub or is_t_sub
        
        c_lat, c_lon = MUNICIPAL_HALL_COORDS.get(c_name, (34.6853, 135.5208))
        cities.append({
            'code': c_code, 'name': c_name, 'wt_j': wt_j, 'wt_t': wt_t,
            'lat': c_lat, 'lon': c_lon,
            'orig_j': off_j_def, 'orig_t': off_t_def,
            'is_sub': is_sub
        })

    for candidate_city, cand_coords in MUNICIPAL_HALL_COORDS.items():
        for j_close in j_existing_list:
            
            active_j_offices = {k[1]: v['coords'] for k, v in EXISTING_OFFICES_INFO.items() if k[0] == 'J社' and k[1] != j_close}
            active_j_offices['J社最適'] = cand_coords
            
            active_t_offices = {k[1]: v['coords'] for k, v in EXISTING_OFFICES_INFO.items() if k[0] == 'T社' and k[1] != t_close_fixed}
            active_t_offices['T社最適'] = cand_coords
            
            total_tk = 0.0
            city_assignments = []
            
            for c_item in cities:
                c_lat, c_lon = c_item['lat'], c_item['lon']
                wt_j_ton = c_item['wt_j'] / 1000.0
                wt_t_ton = c_item['wt_t'] / 1000.0
                
                if c_item['is_sub']:
                    orig_j_clean = c_item['orig_j'].replace('(委託)', '').strip()
                    orig_t_clean = c_item['orig_t'].replace('(委託)', '').strip()
                    
                    j_coords = active_j_offices.get(orig_j_clean, EXISTING_OFFICES_INFO.get(('J社', orig_j_clean), {}).get('coords', (c_lat, c_lon)))
                    t_coords = active_t_offices.get(orig_t_clean, EXISTING_OFFICES_INFO.get(('T社', orig_t_clean), {}).get('coords', (c_lat, c_lon)))
                    
                    dj = calc_haversine_distance(j_coords[0], j_coords[1], c_lat, c_lon) if j_coords else 0.0
                    dt = calc_haversine_distance(t_coords[0], t_coords[1], c_lat, c_lon) if t_coords else 0.0
                    
                    tk_j_c = round(dj * wt_j_ton, 2)
                    tk_t_c = round(dt * wt_t_ton, 2)
                    total_tk += (tk_j_c + tk_t_c)
                    
                    city_assignments.append({
                        '市町村コード': c_item['code'], '市町村名': c_item['name'],
                        '市町村一括担当': 'なし', '市町村一括担当営業所': '-',
                        'J社個別': c_item['orig_j'], 'T社個別': c_item['orig_t'],
                        'J社 距離(km)': dj, 'J社 配達重量(kg)': c_item['wt_j'], 'J社 トンキロ(t・km)': tk_j_c,
                        'T社 距離(km)': dt, 'T社 配達重量(kg)': c_item['wt_t'], 'T社 トンキロ(t・km)': tk_t_c,
                        '_j_def_off': c_item['orig_j'], '_t_def_off': c_item['orig_t']
                    })
                else:
                    best_j_off, best_j_dist = min(
                        ((off, calc_haversine_distance(coords[0], coords[1], c_lat, c_lon)) for off, coords in active_j_offices.items()),
                        key=lambda x: x[1]
                    )
                    best_t_off, best_t_dist = min(
                        ((off, calc_haversine_distance(coords[0], coords[1], c_lat, c_lon)) for off, coords in active_t_offices.items()),
                        key=lambda x: x[1]
                    )
                    
                    dist_bulk_j = calc_haversine_distance(cand_coords[0], cand_coords[1], c_lat, c_lon)
                    tk_bulk_j = dist_bulk_j * (wt_j_ton + wt_t_ton)
                    
                    dist_bulk_t = calc_haversine_distance(cand_coords[0], cand_coords[1], c_lat, c_lon)
                    tk_bulk_t = dist_bulk_t * (wt_j_ton + wt_t_ton)
                    
                    tk_indiv = (best_j_dist * wt_j_ton) + (best_t_dist * wt_t_ton)
                    min_c_tk = min(tk_indiv, tk_bulk_j, tk_bulk_t)
                    
                    if min_c_tk == tk_bulk_j and wt_t_ton > 0:
                        bulk_comp, bulk_off = 'J社', 'J社最適'
                        j_ind, t_ind = '-', '-'
                        dj, dt = dist_bulk_j, dist_bulk_j
                    elif min_c_tk == tk_bulk_t and wt_j_ton > 0:
                        bulk_comp, bulk_off = 'T社', 'T社最適'
                        j_ind, t_ind = '-', '-'
                        dj, dt = dist_bulk_t, dist_bulk_t
                    else:
                        bulk_comp, bulk_off = 'なし', '-'
                        j_ind, t_ind = best_j_off, best_t_off
                        dj, dt = best_j_dist, best_t_dist
                        
                    tk_j_c = round(dj * wt_j_ton, 2)
                    tk_t_c = round(dt * wt_t_ton, 2)
                    total_tk += (tk_j_c + tk_t_c)
                    
                    city_assignments.append({
                        '市町村コード': c_item['code'], '市町村名': c_item['name'],
                        '市町村一括担当': bulk_comp, '市町村一括担当営業所': bulk_off,
                        'J社個別': j_ind, 'T社個別': t_ind,
                        'J社 距離(km)': dj, 'J社 配達重量(kg)': c_item['wt_j'], 'J社 トンキロ(t・km)': tk_j_c,
                        'T社 距離(km)': dt, 'T社 配達重量(kg)': c_item['wt_t'], 'T社 トンキロ(t・km)': tk_t_c,
                        '_j_def_off': c_item['orig_j'], '_t_def_off': c_item['orig_t']
                    })
                    
            if total_tk < best_total_tk:
                best_total_tk = total_tk
                best_config = {
                    'cand_city': candidate_city,
                    'cand_coords': cand_coords,
                    'cand_address': MUNICIPAL_ADDRESS_LOOKUP.get(candidate_city, f"大阪府{candidate_city}"),
                    'j_close': j_close,
                    't_close': t_close_fixed,
                    'total_tk': round(total_tk, 2),
                    'df': pd.DataFrame(city_assignments)
                }

    return best_config

best_opt_res = solve_optimal_scenario_cached(df_j, df_t)

# 変更案設定Excelファイルの取り込み処理
imported_scenario_file = st.sidebar.file_uploader("変更案設定Excelを選択 (案1〜4)", type=["xlsx"], key="scenario_excel_uploader")

if imported_scenario_file is not None:
    try:
        xls_imp = pd.ExcelFile(imported_scenario_file)
        
        if '新規営業所一覧' in xls_imp.sheet_names:
            df_imp_cust = pd.read_excel(xls_imp, sheet_name='新規営業所一覧')
            for _, c_row in df_imp_cust.iterrows():
                p_val = c_row.get('対象パターン', '')
                c_addr = str(c_row.get('所在地住所', '')).strip()
                
                p_nums = parse_pattern_num(p_val)
                if not p_nums:
                    p_nums = [1, 2, 3, 4]
                    
                if c_addr and c_addr != 'nan':
                    coords = geocode_address_precise(c_addr)
                    for p_num in p_nums:
                        j_name_auto = f"J社案{p_num}"
                        t_name_auto = f"T社案{p_num}"
                        
                        existing_names = [x['name'] for x in st.session_state[f"custom_offices_p{p_num}"]]
                        
                        if j_name_auto not in existing_names:
                            st.session_state[f"custom_offices_p{p_num}"].append({
                                'company': 'J社', 'name': j_name_auto, 'address': c_addr, 'coords': coords
                            })
                        if t_name_auto not in existing_names:
                            st.session_state[f"custom_offices_p{p_num}"].append({
                                'company': 'T社', 'name': t_name_auto, 'address': c_addr, 'coords': coords
                            })

        for p_i in [1, 2, 3, 4]:
            sheet_target = f'案{p_i}'
            if sheet_target in xls_imp.sheet_names:
                df_imp_p = pd.read_excel(xls_imp, sheet_name=sheet_target)
                curr_p_df = st.session_state[f"table_data_p{p_i}"]
                
                p_customs = st.session_state[f"custom_offices_p{p_i}"]
                p_off_info = dict(EXISTING_OFFICES_INFO)
                for cust in p_customs:
                    p_off_info[(cust['company'], cust['name'])] = {'address': cust['address'], 'coords': cust['coords']}
                
                for idx, r_imp in df_imp_p.iterrows():
                    c_code = r_imp.get('市町村コード')
                    c_name = str(r_imp.get('市町村名', '')).strip()
                    
                    match_idx = curr_p_df[curr_p_df['市町村名'] == c_name].index
                    if len(match_idx) > 0:
                        i_target = match_idx[0]
                        
                        b_comp = str(r_imp.get('市町村一括担当', curr_p_df.loc[i_target, '市町村一括担当'])).strip()
                        b_off = str(r_imp.get('市町村一括担当営業所', '-')).strip()
                        j_ind = str(r_imp.get('J社個別', curr_p_df.loc[i_target, 'J社個別'])).strip()
                        t_ind = str(r_imp.get('T社個別', curr_p_df.loc[i_target, 'T社個別'])).strip()
                        
                        if b_comp in ['nan', 'None', '', 'なし', '-']:
                            b_comp = 'なし'
                            b_off = '-'
                        if b_off in ['nan', 'None', '']:
                            b_off = '-'
                        if j_ind in ['nan', 'None', '']:
                            j_ind = '-'
                        if t_ind in ['nan', 'None', '']:
                            t_ind = '-'
                        
                        wt_j = curr_p_df.loc[i_target, 'J社 配達重量(kg)']
                        wt_t = curr_p_df.loc[i_target, 'T社 配達重量(kg)']
                        
                        dj, tkj, dt, tkt = compute_distance_and_tonkm(c_name, b_comp, b_off, j_ind, t_ind, wt_j, wt_t, p_off_info, MUNICIPAL_HALL_COORDS)
                        
                        curr_p_df.at[i_target, '市町村一括担当'] = b_comp
                        curr_p_df.at[i_target, '市町村一括担当営業所'] = b_off
                        curr_p_df.at[i_target, 'J社個別'] = j_ind
                        curr_p_df.at[i_target, 'T社個別'] = t_ind
                        curr_p_df.at[i_target, 'J社 距離(km)'] = dj
                        curr_p_df.at[i_target, 'J社 トンキロ(t・km)'] = tkj
                        curr_p_df.at[i_target, 'T社 距離(km)'] = dt
                        curr_p_df.at[i_target, 'T社 トンキロ(t・km)'] = tkt
                        
                st.session_state[f"table_data_p{p_i}"] = curr_p_df

        st.sidebar.success("変更案設定Excel（案1〜4）を取り込みました！")
    except Exception as e_imp:
        st.sidebar.error(f"Excel取り込みエラー: {e_imp}")

# 5. タブUI構成
tab_names = [
    "📊 全体サマリー＆現状比較",
    "🏢 営業所別 集計比較",
    "📝 案1：受持選択（編集）",
    "📝 案2：受持選択（編集）",
    "📝 案3：受持選択（編集）",
    "📝 案4：受持選択（編集）",
    "⭐ 最適解：受持選択"
]
tabs = st.tabs(tab_names)

pattern_active_records = {1: [], 2: [], 3: [], 4: []}
pattern_map_status = {1: {}, 2: {}, 3: {}, 4: {}}
pattern_map_images = {1: None, 2: None, 3: None, 4: None}

display_cols = [
    '市町村コード', '市町村名', '市町村一括担当', '市町村一括担当営業所',
    'J社個別', 'T社個別',
    'J社 距離(km)', 'J社 配達重量(kg)', 'J社 トンキロ(t・km)',
    'T社 距離(km)', 'T社 配達重量(kg)', 'T社 トンキロ(t・km)'
]

# タブ3〜6: 案1〜案4 個別編集画面処理
for p_idx in [1, 2, 3, 4]:
    tab_obj = tabs[p_idx + 1]
    with tab_obj:
        st.subheader(f"📝 案{p_idx}：市町村別・受持選択テーブル")
        st.caption(f"※ 案{p_idx} の設定を行います。所在地住所を入力して登録すると、J社（J社案{p_idx}）とT社（T社案{p_idx}）の新規営業所が一括登録されます。")
        
        p_custom_offices = st.session_state[f"custom_offices_p{p_idx}"]
        
        j_color_map_p, t_color_map_p, j_offices_p, t_offices_p = generate_office_colors(df_j, df_t, p_custom_offices)
        all_offices_p = sorted(list(set(j_offices_p + t_offices_p)))
        
        p_office_info = dict(EXISTING_OFFICES_INFO)
        for cust in p_custom_offices:
            p_office_info[(cust['company'], cust['name'])] = {'address': cust['address'], 'coords': cust['coords']}
        
        st.markdown(f"##### ➕ 案{p_idx} 専用 新規営業所の登録（J社案{p_idx} / T社案{p_idx} 一括登録）")
        col_addr_shared, col_btn_shared = st.columns([3, 1])
        with col_addr_shared:
            new_shared_addr = st.text_input("新規営業所 所在地住所 (J社・T社共通)", key=f"shared_off_addr_p{p_idx}", placeholder="例: 大阪府茨木市駅前1-1-1")
        with col_btn_shared:
            st.write("")
            st.write("")
            if st.button(f"J社案{p_idx} / T社案{p_idx} を追加", key=f"btn_add_shared_p{p_idx}"):
                if new_shared_addr:
                    detected_coords = geocode_address_precise(new_shared_addr)
                    
                    j_name_auto = f"J社案{p_idx}"
                    t_name_auto = f"T社案{p_idx}"
                    
                    existing_names = [x['name'] for x in st.session_state[f"custom_offices_p{p_idx}"]]
                    
                    if j_name_auto not in existing_names:
                        st.session_state[f"custom_offices_p{p_idx}"].append({
                            'company': 'J社', 'name': j_name_auto, 'address': new_shared_addr, 'coords': detected_coords
                        })
                    if t_name_auto not in existing_names:
                        st.session_state[f"custom_offices_p{p_idx}"].append({
                            'company': 'T社', 'name': t_name_auto, 'address': new_shared_addr, 'coords': detected_coords
                        })
                        
                    st.success(f"案{p_idx} に 「{j_name_auto}」および「{t_name_auto}」（住所: {new_shared_addr}）を一括追加しました！")
                    st.rerun()
        st.markdown("---")

        curr_p_df = st.session_state[f"table_data_p{p_idx}"]
        
        recalc_clicked = st.button(f"🔄 案{p_idx} の配達支店変更を反映", type="primary", key=f"btn_recalc_p{p_idx}", use_container_width=True)

        edited_p_df = st.data_editor(
            curr_p_df[display_cols],
            column_config={
                "市町村一括担当": st.column_config.SelectboxColumn("市町村一括担当", options=["なし", "J社", "T社"], required=True),
                "市町村一括担当営業所": st.column_config.SelectboxColumn("市町村一括担当営業所", options=all_offices_p + ["-"], required=True),
                "J社個別": st.column_config.SelectboxColumn("J社個別", options=j_offices_p + ["-"], required=True),
                "T社個別": st.column_config.SelectboxColumn("T社個別", options=t_offices_p + ["-"], required=True),
                "J社 距離(km)": st.column_config.NumberColumn("J社 距離 (km)", format="%.1f"),
                "J社 配達重量(kg)": st.column_config.NumberColumn("J社 配達重量 (kg)", format="%d"),
                "J社 トンキロ(t・km)": st.column_config.NumberColumn("J社 トンキロ (t・km)", format="%.2f"),
                "T社 距離(km)": st.column_config.NumberColumn("T社 距離 (km)", format="%.1f"),
                "T社 配達重量(kg)": st.column_config.NumberColumn("T社 配達重量 (kg)", format="%d"),
                "T社 トンキロ(t・km)": st.column_config.NumberColumn("T社 トンキロ (t・km)", format="%.2f"),
            },
            disabled=['市町村コード', '市町村名', 'J社 距離(km)', 'J社 配達重量(kg)', 'J社 トンキロ(t・km)', 'T社 距離(km)', 'T社 配達重量(kg)', 'T社 トンキロ(t・km)'],
            use_container_width=True,
            hide_index=True,
            key=f"main_editor_p{p_idx}"
        )

        if recalc_clicked:
            for idx in range(len(edited_p_df)):
                r = edited_p_df.iloc[idx]
                c_name = r['市町村名']
                b_comp = r['市町村一括担当']
                b_off = r['市町村一括担当営業所']
                j_ind = r['J社個別']
                t_ind = r['T社個別']
                wt_j = r['J社 配達重量(kg)']
                wt_t = r['T社 配達重量(kg)']
                
                dj, tkj, dt, tkt = compute_distance_and_tonkm(c_name, b_comp, b_off, j_ind, t_ind, wt_j, wt_t, p_office_info, MUNICIPAL_HALL_COORDS)
                
                curr_p_df.at[idx, '市町村一括担当'] = b_comp
                curr_p_df.at[idx, '市町村一括担当営業所'] = b_off
                curr_p_df.at[idx, 'J社個別'] = j_ind
                curr_p_df.at[idx, 'T社個別'] = t_ind
                curr_p_df.at[idx, 'J社 距離(km)'] = dj
                curr_p_df.at[idx, 'J社 トンキロ(t・km)'] = tkj
                curr_p_df.at[idx, 'T社 距離(km)'] = dt
                curr_p_df.at[idx, 'T社 トンキロ(t・km)'] = tkt
                
            st.session_state[f"table_data_p{p_idx}"] = curr_p_df
            st.success(f"案{p_idx} の距離およびトンキロを最新表示に再計算しました！")
            st.rerun()

        p_active = []
        p_map_dict = {}
        for idx, r in edited_p_df.iterrows():
            orig = curr_p_df.loc[idx]
            city_name = orig['市町村名']
            city_code = orig['市町村コード']
            
            bulk_comp = r['市町村一括担当']
            bulk_off = r['市町村一括担当営業所']
            j_indiv_off = r['J社個別']
            t_indiv_off = r['T社個別']
            
            dj_now, tkj_now, dt_now, tkt_now = compute_distance_and_tonkm(city_name, bulk_comp, bulk_off, j_indiv_off, t_indiv_off, orig['J社 配達重量(kg)'], orig['T社 配達重量(kg)'], p_office_info, MUNICIPAL_HALL_COORDS)
            
            curr_p_df.at[idx, '市町村一括担当'] = bulk_comp
            curr_p_df.at[idx, '市町村一括担当営業所'] = bulk_off
            curr_p_df.at[idx, 'J社個別'] = j_indiv_off
            curr_p_df.at[idx, 'T社個別'] = t_indiv_off
            curr_p_df.at[idx, 'J社 距離(km)'] = dj_now
            curr_p_df.at[idx, 'J社 トンキロ(t・km)'] = tkj_now
            curr_p_df.at[idx, 'T社 距離(km)'] = dt_now
            curr_p_df.at[idx, 'T社 トンキロ(t・km)'] = tkt_now

            dist_j = dj_now
            dist_t = dt_now
            tk_j = tkj_now
            tk_t = tkt_now
            
            wt_j_kg = orig['J社 配達重量(kg)']
            wt_t_kg = orig['T社 配達重量(kg)']
            wt_j_t = wt_j_kg / 1000.0
            wt_t_t = wt_t_kg / 1000.0
            
            p_map_dict[city_name] = {
                '一括担当': bulk_comp,
                '一括営業所': bulk_off if bulk_off != '-' else (j_indiv_off if bulk_comp == 'J社' and j_indiv_off != '-' else (t_indiv_off if bulk_comp == 'T社' and t_indiv_off != '-' else orig['_j_def_off'] if bulk_comp == 'J社' else orig['_t_def_off'])),
                'j_indiv_off': j_indiv_off,
                't_indiv_off': t_indiv_off
            }
            
            has_indiv = (j_indiv_off != '-') or (t_indiv_off != '-')
            if has_indiv:
                if j_indiv_off != '-':
                    p_active.append({
                        '市区町村コード': city_code, '市区町村名': city_name,
                        '担当会社': 'J社', '担当営業所': j_indiv_off,
                        '重量_t': wt_j_t, '距離_km': dist_j, 'トンキロ': tk_j,
                    })
                if t_indiv_off != '-':
                    p_active.append({
                        '市区町村コード': city_code, '市区町村名': city_name,
                        '担当会社': 'T社', '担当営業所': t_indiv_off,
                        '重量_t': wt_t_t, '距離_km': dist_t, 'トンキロ': tk_t,
                    })
            else:
                if bulk_comp in ['J社', 'T社']:
                    assigned_off = bulk_off if bulk_off != '-' else (orig['_j_def_off'] if bulk_comp == 'J社' else orig['_t_def_off'])
                    c_code = bulk_comp[0]
                    total_wt = wt_j_t + wt_t_t
                    avg_dist = dist_j if c_code == 'J' else dist_t
                    total_tk = round(avg_dist * total_wt, 2)
                    
                    p_active.append({
                        '市区町村コード': city_code, '市区町村名': city_name,
                        '担当会社': bulk_comp, '担当営業所': assigned_off,
                        '重量_t': total_wt, '距離_km': avg_dist, 'トンキロ': total_tk,
                    })

        st.session_state[f"table_data_p{p_idx}"] = curr_p_df
        pattern_active_records[p_idx] = p_active
        pattern_map_status[p_idx] = p_map_dict

        st.markdown("---")
        st.subheader(f"🗺️ 案{p_idx} 白地図エリアマップ")
        
        st.markdown("##### 📌 営業所別 カラー凡例（拡大カラーブロック表示）")
        leg_cols = st.columns(2)
        with leg_cols[0]:
            st.write("**🔴 J社 営業所**")
            for off in j_offices_p:
                c = j_color_map_p.get(off, (239, 68, 68))
                hex_c = f"#{c[0]:02x}{c[1]:02x}{c[2]:02x}"
                st.markdown(f'<span style="color:{hex_c}; font-size:48px; vertical-align:middle; line-height:1;">■</span> <span style="font-size:18px; font-weight:bold; vertical-align:middle;">{off}</span>', unsafe_allow_html=True)
                
        with leg_cols[1]:
            st.write("**🔵 T社 営業所**")
            for off in t_offices_p:
                c = t_color_map_p.get(off, (59, 130, 246))
                hex_c = f"#{c[0]:02x}{c[1]:02x}{c[2]:02x}"
                st.markdown(f'<span style="color:{hex_c}; font-size:48px; vertical-align:middle; line-height:1;">■</span> <span style="font-size:18px; font-weight:bold; vertical-align:middle;">{off}</span>', unsafe_allow_html=True)

        st.write("")

        map_img_path = None
        for target_path in ["画像4_2.png", "画像4_2.jpg", "画像4.png", "画像3.jpg"]:
            if os.path.exists(target_path):
                map_img_path = target_path
                break

        if map_img_path and os.path.exists(map_img_path):
            src_img = cv2.imread(map_img_path)
            src_img = cv2.resize(src_img, (1024, 1400), interpolation=cv2.INTER_AREA)
            h, w, _ = src_img.shape
            
            gray = cv2.cvtColor(src_img, cv2.COLOR_BGR2GRAY)
            line_bin = (gray > 40).astype(np.uint8) * 255
                
            kernel = np.ones((3, 3), np.uint8)
            line_bin_dilated = cv2.dilate(line_bin, kernel, iterations=1)
            
            canvas_rgb = np.full((h, w, 3), 255, dtype=np.uint8)
            canvas_rgb[line_bin_dilated == 255] = (0, 0, 0)
            
            bg_protection_mask = np.zeros((h + 2, w + 2), np.uint8)
            cv2.floodFill(canvas_rgb.copy(), bg_protection_mask, (0, 0), (255, 255, 255), (15, 15, 15), (15, 15, 15), cv2.FLOODFILL_FIXED_RANGE)

            for c_name, (sx, sy) in CITY_SEEDS.items():
                info = p_map_dict.get(c_name, {'一括担当': 'なし', '一括営業所': '-', 'j_indiv_off': '-', 't_indiv_off': '-'})
                bulk = info['一括担当']
                bulk_off = info['一括営業所']
                j_indiv_off = info.get('j_indiv_off', '-')
                t_indiv_off = info.get('t_indiv_off', '-')
                
                if bulk == 'J社' and bulk_off in j_color_map_p:
                    fill_rgb = j_color_map_p[bulk_off]
                elif bulk == 'T社' and bulk_off in t_color_map_p:
                    fill_rgb = t_color_map_p[bulk_off]
                elif bulk in ['なし', '-'] and ((j_indiv_off != '-') or (t_indiv_off != '-')):
                    fill_rgb = (200, 200, 200) # 個別選択あり＝灰色
                else:
                    fill_rgb = (255, 255, 255) # 未設定＝白色

                target_coords = [(sx, sy)]
                if c_name in EXTRA_SEEDS:
                    target_coords.extend(EXTRA_SEEDS[c_name])
                    
                for cs_x, cs_y in target_coords:
                    # 安全境界ガード（IndexError完全回避）
                    cs_x = int(cs_x)
                    cs_y = int(cs_y)
                    if 0 <= cs_x < w and 0 <= cs_y < h:
                        if line_bin_dilated[cs_y, cs_x] == 255:
                            found = False
                            for r in range(1, 15):
                                for dy in range(-r, r+1):
                                    for dx in range(-r, r+1):
                                        nx, ny = cs_x + dx, cs_y + dy
                                        if 0 <= nx < w and 0 <= ny < h and line_bin_dilated[ny, nx] == 0:
                                            cs_x, cs_y = nx, ny
                                            found = True
                                            break
                                    if found: break
                                if found: break

                        if 0 <= cs_x < w and 0 <= cs_y < h and line_bin_dilated[cs_y, cs_x] == 0:
                            m_curr = bg_protection_mask.copy()
                            cv2.floodFill(canvas_rgb, m_curr, (cs_x, cs_y), fill_rgb, (15, 15, 15), (15, 15, 15), cv2.FLOODFILL_FIXED_RANGE)

            canvas_rgb[line_bin_dilated == 255] = (0, 0, 0)

            gray_check = cv2.cvtColor(canvas_rgb, cv2.COLOR_RGB2GRAY)
            non_bg = np.where(gray_check < 250)
            
            if len(non_bg[0]) > 0:
                min_y, max_y = np.min(non_bg[0]), np.max(non_bg[0])
                min_x, max_x = np.min(non_bg[1]), np.max(non_bg[1])
                pad = 12
                crop_min_y = max(0, min_y - pad)
                crop_max_y = min(h, max_y + pad)
                crop_min_x = max(0, min_x - pad)
                crop_max_x = min(w, max_x + pad)
                cropped_img = canvas_rgb[crop_min_y:crop_max_y, crop_min_x:crop_max_x]
            else:
                cropped_img = canvas_rgb

            pattern_map_images[p_idx] = cropped_img

            map_col1, map_col2, map_col3 = st.columns([1, 4, 1])
            with map_col2:
                st.image(cropped_img, width=600)

# ③ UIタブ「⭐ 最適解：受持選択」のレンダリング（白地図エリアマップ画像付き ＆ 境界ガード保護）
with tabs[6]:
    st.subheader("⭐ トンキロ最小化：全体最適解（1拠点新設 ＆ 既存2拠点廃止）")
    st.caption("数学的最適化アルゴリズムにより、全体トンキロが最も小さくなる「新規共同拠点の1箇所設置」および「J社・T社既存拠点の各1箇所廃止」の最小化構成を全自動計算した結果です。")
    
    col_opt_info1, col_opt_info2, col_opt_info3 = st.columns(3)
    with col_opt_info1:
        st.info(f"📍 **新規営業所 設置地 (J社/T社共同):**\n\n**{best_opt_res['cand_city']}**（{best_opt_res['cand_address']}）")
    with col_opt_info2:
        st.warning(f"廃止対象営業所:\n- 🔴 **J社 廃止拠点:** {best_opt_res['j_close']}\n- 🔵 **T社 廃止拠点 (固定):** {best_opt_res['t_close']}")
    with col_opt_info3:
        base_0_df_opt = build_initial_scenario_df(EXISTING_OFFICES_INFO)
        cur_tk_tot_opt = base_0_df_opt['J社 トンキロ(t・km)'].sum() + base_0_df_opt['T社 トンキロ(t・km)'].sum()
        opt_tk_tot = best_opt_res['total_tk']
        diff_opt = cur_tk_tot_opt - opt_tk_tot
        pct_opt = round((diff_opt / cur_tk_tot_opt * 100), 1) if cur_tk_tot_opt > 0 else 0.0
        st.success(f"🏆 **最適化合計トンキロ:**\n\n**{opt_tk_tot:,.2f} ton-km**\n\n(現状比 削減: {diff_opt:,.1f} ton-km / {pct_opt:+.1f}%)")

    st.markdown("---")
    st.markdown("#### 📋 最適解 市区町村別 受持割り当て一覧")
    st.dataframe(
        best_opt_res['df'][display_cols],
        column_config={
            "J社 距離(km)": st.column_config.NumberColumn("J社 距離 (km)", format="%.1f"),
            "J社 配達重量(kg)": st.column_config.NumberColumn("J社 配達重量 (kg)", format="%d"),
            "J社 トンキロ(t・km)": st.column_config.NumberColumn("J社 トンキロ (t・km)", format="%.2f"),
            "T社 距離(km)": st.column_config.NumberColumn("T社 距離 (km)", format="%.1f"),
            "T社 配達重量(kg)": st.column_config.NumberColumn("T社 配達重量 (kg)", format="%d"),
            "T社 トンキロ(t・km)": st.column_config.NumberColumn("T社 トンキロ (t・km)", format="%.2f"),
        },
        use_container_width=True,
        hide_index=True
    )

    # ③ 最適解の白地図エリアマップ描画（完全境界ガード保護）
    st.markdown("---")
    st.subheader("🗺️ 最適解 白地図エリアマップ")
    opt_custom_offices = [
        {'company': 'J社', 'name': 'J社最適', 'address': best_opt_res['cand_address'], 'coords': best_opt_res['cand_coords']},
        {'company': 'T社', 'name': 'T社最適', 'address': best_opt_res['cand_address'], 'coords': best_opt_res['cand_coords']}
    ]
    j_cmap_opt, t_cmap_opt, j_offs_opt, t_offs_opt = generate_office_colors(df_j, df_t, opt_custom_offices)
    
    st.markdown("##### 📌 営業所別 カラー凡例（拡大カラーブロック表示）")
    leg_cols_opt = st.columns(2)
    with leg_cols_opt[0]:
        st.write("**🔴 J社 営業所**")
        for off in j_offs_opt:
            c = j_cmap_opt.get(off, (239, 68, 68))
            hex_c = f"#{c[0]:02x}{c[1]:02x}{c[2]:02x}"
            st.markdown(f'<span style="color:{hex_c}; font-size:48px; vertical-align:middle; line-height:1;">■</span> <span style="font-size:18px; font-weight:bold; vertical-align:middle;">{off}</span>', unsafe_allow_html=True)
            
    with leg_cols_opt[1]:
        st.write("**🔵 T社 営業所**")
        for off in t_offs_opt:
            c = t_cmap_opt.get(off, (59, 130, 246))
            hex_c = f"#{c[0]:02x}{c[1]:02x}{c[2]:02x}"
            st.markdown(f'<span style="color:{hex_c}; font-size:48px; vertical-align:middle; line-height:1;">■</span> <span style="font-size:18px; font-weight:bold; vertical-align:middle;">{off}</span>', unsafe_allow_html=True)

    st.write("")

    map_img_path_opt = None
    for target_path in ["画像4_2.png", "画像4_2.jpg", "画像4.png", "画像3.jpg"]:
        if os.path.exists(target_path):
            map_img_path_opt = target_path
            break

    opt_map_image = None
    if map_img_path_opt and os.path.exists(map_img_path_opt):
        src_img_o = cv2.imread(map_img_path_opt)
        src_img_o = cv2.resize(src_img_o, (1024, 1400), interpolation=cv2.INTER_AREA)
        h_o, w_o, _ = src_img_o.shape
        
        gray_o = cv2.cvtColor(src_img_o, cv2.COLOR_BGR2GRAY)
        line_bin_o = (gray_o > 40).astype(np.uint8) * 255
            
        kernel_o = np.ones((3, 3), np.uint8)
        line_bin_dilated_o = cv2.dilate(line_bin_o, kernel_o, iterations=1)
        
        canvas_rgb_o = np.full((h_o, w_o, 3), 255, dtype=np.uint8)
        canvas_rgb_o[line_bin_dilated_o == 255] = (0, 0, 0)
        
        bg_protection_mask_o = np.zeros((h_o + 2, w_o + 2), np.uint8)
        cv2.floodFill(canvas_rgb_o.copy(), bg_protection_mask_o, (0, 0), (255, 255, 255), (15, 15, 15), (15, 15, 15), cv2.FLOODFILL_FIXED_RANGE)

        df_opt_map = best_opt_res['df']
        for _, r_opt_m in df_opt_map.iterrows():
            c_name_m = r_opt_m['市町村名']
            bulk_m = r_opt_m['市町村一括担当']
            bulk_off_m = r_opt_m['市町村一括担当営業所']
            j_ind_m = r_opt_m['J社個別']
            t_ind_m = r_opt_m['T社個別']
            
            if bulk_m == 'J社' and bulk_off_m in j_cmap_opt:
                fill_rgb_o = j_cmap_opt[bulk_off_m]
            elif bulk_m == 'T社' and bulk_off_m in t_cmap_opt:
                fill_rgb_o = t_cmap_opt[bulk_off_m]
            elif bulk_m in ['なし', '-'] and ((j_ind_m != '-') or (t_ind_m != '-')):
                fill_rgb_o = (200, 200, 200)
            else:
                fill_rgb_o = (255, 255, 255)

            if c_name_m in CITY_SEEDS:
                sx_o, sy_o = CITY_SEEDS[c_name_m]
                target_coords_o = [(sx_o, sy_o)]
                if c_name_m in EXTRA_SEEDS:
                    target_coords_o.extend(EXTRA_SEEDS[c_name_m])
                    
                for cs_x_o, cs_y_o in target_coords_o:
                    cs_x_o = int(cs_x_o)
                    cs_y_o = int(cs_y_o)
                    # 完全安全境界ガード
                    if 0 <= cs_x_o < w_o and 0 <= cs_y_o < h_o:
                        if line_bin_dilated_o[cs_y_o, cs_x_o] == 255:
                            found_o = False
                            for r_o in range(1, 15):
                                for dy_o in range(-r_o, r_o+1):
                                    for dx_o in range(-r_o, r_o+1):
                                        nx_o, ny_o = cs_x_o + dx_o, cs_y_o + dy_o
                                        if 0 <= nx_o < w_o and 0 <= ny_o < h_o and line_bin_dilated_o[ny_o, nx_o] == 0:
                                            cs_x_o, cs_y_o = nx_o, ny_o
                                            found_o = True
                                            break
                                    if found_o: break
                                if found_o: break

                        if 0 <= cs_x_o < w_o and 0 <= cs_y_o < h_o and line_bin_dilated_o[cs_y_o, cs_x_o] == 0:
                            m_curr_o = bg_protection_mask_o.copy()
                            cv2.floodFill(canvas_rgb_o, m_curr_o, (cs_x_o, cs_y_o), fill_rgb_o, (15, 15, 15), (15, 15, 15), cv2.FLOODFILL_FIXED_RANGE)

        canvas_rgb_o[line_bin_dilated_o == 255] = (0, 0, 0)
        gray_check_o = cv2.cvtColor(canvas_rgb_o, cv2.COLOR_RGB2GRAY)
        non_bg_o = np.where(gray_check_o < 250)
        
        if len(non_bg_o[0]) > 0:
            min_y_o, max_y_o = np.min(non_bg_o[0]), np.max(non_bg_o[0])
            min_x_o, max_x_o = np.min(non_bg_o[1]), np.max(non_bg_o[1])
            pad_o = 12
            crop_min_y_o = max(0, min_y_o - pad_o)
            crop_max_y_o = min(h_o, max_y_o + pad_o)
            crop_min_x_o = max(0, min_x_o - pad_o)
            crop_max_x_o = min(w_o, max_x_o + pad_o)
            opt_map_image = canvas_rgb_o[crop_min_y_o:crop_max_y_o, crop_min_x_o:crop_max_x_o]
        else:
            opt_map_image = canvas_rgb_o

        map_col1_o, map_col2_o, map_col3_o = st.columns([1, 4, 1])
        with map_col2_o:
            st.image(opt_map_image, width=600)

# タブ1: 全体サマリー ＆ 現状 vs 4つの変更案 ＋ 最適解比較
with tabs[0]:
    st.subheader("📈 会社毎の「現状」vs「4つの変更案（案1〜案4）」vs「最適解」一括比較サマリー")
    st.caption("J社・T社個別の数値に加え、運送2社全体の合計トンキロ・配達重量（t）・削減率を一括比較します。")

    base_0_df = build_initial_scenario_df(EXISTING_OFFICES_INFO)
    
    cur_j_wt_total = base_0_df['J社 配達重量(kg)'].sum() / 1000.0
    cur_t_wt_total = base_0_df['T社 配達重量(kg)'].sum() / 1000.0
    cur_j_tk_total = base_0_df['J社 トンキロ(t・km)'].sum()
    cur_t_tk_total = base_0_df['T社 トンキロ(t・km)'].sum()

    summary_rows = [
        {'パターン': '現状（基本）', '会社': 'J社', '担当自治体数': len(base_0_df), '配達重量_t': round(cur_j_wt_total, 2), 'トンキロ': round(cur_j_tk_total, 2), 'トンキロ削減量': 0.0, '削減率(%)': 0.0},
        {'パターン': '現状（基本）', '会社': 'T社', '担当自治体数': len(base_0_df), '配達重量_t': round(cur_t_wt_total, 2), 'トンキロ': round(cur_t_tk_total, 2), 'トンキロ削減量': 0.0, '削減率(%)': 0.0}
    ]

    for p_idx in [1, 2, 3, 4]:
        p_df = st.session_state[f"table_data_p{p_idx}"]
        
        j_tk_sum = 0.0; t_tk_sum = 0.0
        j_wt_sum = 0.0; t_wt_sum = 0.0
        j_cnt = 0; t_cnt = 0
        
        for _, r in p_df.iterrows():
            bulk = str(r['市町村一括担当']).strip()
            j_tk = float(r['J社 トンキロ(t・km)'])
            t_tk = float(r['T社 トンキロ(t・km)'])
            j_wt = float(r['J社 配達重量(kg)']) / 1000.0
            t_wt = float(r['T社 配達重量(kg)']) / 1000.0
            
            if bulk == 'J社':
                j_tk_sum += (j_tk + t_tk)
                j_wt_sum += (j_wt + t_wt)
                j_cnt += 1
            elif bulk == 'T社':
                t_tk_sum += (j_tk + t_tk)
                t_wt_sum += (j_wt + t_wt)
                t_cnt += 1
            else:
                j_tk_sum += j_tk
                j_wt_sum += j_wt
                if j_tk > 0 or str(r['J社個別']).strip() != '-':
                    j_cnt += 1
                    
                t_tk_sum += t_tk
                t_wt_sum += t_wt
                if t_tk > 0 or str(r['T社個別']).strip() != '-':
                    t_cnt += 1

        j_diff = cur_j_tk_total - j_tk_sum
        j_pct = round((j_diff / cur_j_tk_total * 100), 1) if cur_j_tk_total > 0 else 0.0
        summary_rows.append({'パターン': f'案{p_idx}', '会社': 'J社', '担当自治体数': j_cnt, '配達重量_t': round(j_wt_sum, 2), 'トンキロ': round(j_tk_sum, 2), 'トンキロ削減量': round(j_diff, 2), '削減率(%)': j_pct})

        t_diff = cur_t_tk_total - t_tk_sum
        t_pct = round((t_diff / cur_t_tk_total * 100), 1) if cur_t_tk_total > 0 else 0.0
        summary_rows.append({'パターン': f'案{p_idx}', '会社': 'T社', '担当自治体数': t_cnt, '配達重量_t': round(t_wt_sum, 2), 'トンキロ': round(t_tk_sum, 2), 'トンキロ削減量': round(t_diff, 2), '削減率(%)': t_pct})

    df_opt_calc = best_opt_res['df']
    opt_j_tk_sum = df_opt_calc['J社 トンキロ(t・km)'].sum()
    opt_t_tk_sum = df_opt_calc['T社 トンキロ(t・km)'].sum()
    opt_j_wt_sum = df_opt_calc['J社 配達重量(kg)'].sum() / 1000.0
    opt_t_wt_sum = df_opt_calc['T社 配達重量(kg)'].sum() / 1000.0
    opt_j_cnt = len(df_opt_calc[df_opt_calc['J社 トンキロ(t・km)'] > 0])
    opt_t_cnt = len(df_opt_calc[df_opt_calc['T社 トンキロ(t・km)'] > 0])

    opt_j_diff = cur_j_tk_total - opt_j_tk_sum
    opt_j_pct = round((opt_j_diff / cur_j_tk_total * 100), 1) if cur_j_tk_total > 0 else 0.0
    summary_rows.append({'パターン': '最適解', '会社': 'J社', '担当自治体数': opt_j_cnt, '配達重量_t': round(opt_j_wt_sum, 2), 'トンキロ': round(opt_j_tk_sum, 2), 'トンキロ削減量': round(opt_j_diff, 2), '削減率(%)': opt_j_pct})

    opt_t_diff = cur_t_tk_total - opt_t_tk_sum
    opt_t_pct = round((opt_t_diff / cur_t_tk_total * 100), 1) if cur_t_tk_total > 0 else 0.0
    summary_rows.append({'パターン': '最適解', '会社': 'T社', '担当自治体数': opt_t_cnt, '配達重量_t': round(opt_t_wt_sum, 2), 'トンキロ': round(opt_t_tk_sum, 2), 'トンキロ削減量': round(opt_t_diff, 2), '削減率(%)': opt_t_pct})

    summary_df = pd.DataFrame(summary_rows)

    cur_comb_tk_total = cur_j_tk_total + cur_t_tk_total
    cur_comb_wt_total = cur_j_wt_total + cur_t_wt_total
    
    comb_summary_rows = []
    
    comb_summary_rows.append({
        'パターン': '現状（基本）',
        'J社_トンキロ': round(cur_j_tk_total, 2),
        'T社_トンキロ': round(cur_t_tk_total, 2),
        '合計_トンキロ': round(cur_comb_tk_total, 2),
        '合計_削減量': 0.0,
        '合計_削減率(%)': 0.0,
    })

    for p_idx in [1, 2, 3, 4]:
        j_row = summary_df[(summary_df['パターン'] == f'案{p_idx}') & (summary_df['会社'] == 'J社')].iloc[0]
        t_row = summary_df[(summary_df['パターン'] == f'案{p_idx}') & (summary_df['会社'] == 'T社')].iloc[0]
        
        j_tk_p = j_row['トンキロ']
        t_tk_p = t_row['トンキロ']
        comb_tk_p = j_tk_p + t_tk_p
        comb_diff = cur_comb_tk_total - comb_tk_p
        comb_pct = round((comb_diff / cur_comb_tk_total * 100), 1) if cur_comb_tk_total > 0 else 0.0
        
        comb_summary_rows.append({
            'パターン': f'案{p_idx}',
            'J社_トンキロ': j_tk_p,
            'T社_トンキロ': t_tk_p,
            '合計_トンキロ': round(comb_tk_p, 2),
            '合計_削減量': round(comb_diff, 2),
            '合計_削減率(%)': comb_pct,
        })

    opt_comb_tk_p = opt_j_tk_sum + opt_t_tk_sum
    opt_comb_diff = cur_comb_tk_total - opt_comb_tk_p
    opt_comb_pct = round((opt_comb_diff / cur_comb_tk_total * 100), 1) if cur_comb_tk_total > 0 else 0.0
    comb_summary_rows.append({
        'パターン': '最適解',
        'J社_トンキロ': round(opt_j_tk_sum, 2),
        'T社_トンキロ': round(opt_t_tk_sum, 2),
        '合計_トンキロ': round(opt_comb_tk_p, 2),
        '合計_削減量': round(opt_comb_diff, 2),
        '合計_削減率(%)': opt_comb_pct,
    })

    comb_summary_df = pd.DataFrame(comb_summary_rows)

    excel_full_bytes = io.BytesIO()
    with pd.ExcelWriter(excel_full_bytes, engine='openpyxl') as writer:
        all_cust_rows = []
        for p_i in [1, 2, 3, 4]:
            p_custs = st.session_state[f"custom_offices_p{p_i}"]
            for c_item in p_custs:
                all_cust_rows.append({
                    '対象パターン': f'案{p_i}',
                    '担当会社': c_item['company'],
                    '新規営業所名': c_item['name'],
                    '所在地住所': c_item['address']
                })
                
        all_cust_rows.append({
            '対象パターン': '最適解',
            '担当会社': 'J社',
            '新規営業所名': 'J社最適',
            '所在地住所': best_opt_res['cand_address']
        })
        all_cust_rows.append({
            '対象パターン': '最適解',
            '担当会社': 'T社',
            '新規営業所名': 'T社最適',
            '所在地住所': best_opt_res['cand_address']
        })

        df_cust_all = pd.DataFrame(all_cust_rows)
        df_cust_all.to_excel(writer, sheet_name='新規営業所一覧', index=False)

        df_cur_export = base_0_df[display_cols].copy()
        df_cur_export.to_excel(writer, sheet_name='現状', index=False)

        for p_i in [1, 2, 3, 4]:
            p_df_export = st.session_state[f"table_data_p{p_i}"][display_cols].copy()
            p_df_export.to_excel(writer, sheet_name=f'案{p_i}', index=False)

        df_opt_export = best_opt_res['df'][display_cols].copy()
        df_opt_export.to_excel(writer, sheet_name='最適解', index=False)

        wb = writer.book

        hdr_font = Font(name='BIZ UDPゴシック', size=11, bold=False, color='FFFFFF')
        hdr_fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
        data_font = Font(name='BIZ UDPゴシック', size=11, bold=False)
        border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

        for sheet_name in ['新規営業所一覧', '現状', '案1', '案2', '案3', '案4', '最適解']:
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                ws.views.sheetView[0].showGridLines = True
                
                for cell in ws[1]:
                    cell.font = hdr_font
                    cell.fill = hdr_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = border_thin
                
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.font = data_font
                        cell.border = border_thin

        ws_cust = wb['新規営業所一覧']
        ws_cust.column_dimensions['A'].width = 13.5
        ws_cust.column_dimensions['B'].width = 9.6
        ws_cust.column_dimensions['C'].width = 13.9
        ws_cust.column_dimensions['D'].width = 37.5

        for sheet_item in ['現状', '案1', '案2', '案3', '案4', '最適解']:
            if sheet_item in wb.sheetnames:
                ws_p = wb[sheet_item]
                ws_p.column_dimensions['A'].width = 12.5
                ws_p.column_dimensions['B'].width = 16.1
                ws_p.column_dimensions['C'].width = 14.0
                ws_p.column_dimensions['D'].width = 22.2
                ws_p.column_dimensions['E'].width = 12.7
                ws_p.column_dimensions['F'].width = 12.7
                ws_p.column_dimensions['G'].width = 14.9
                ws_p.column_dimensions['H'].width = 18.6
                ws_p.column_dimensions['I'].width = 20.4
                ws_p.column_dimensions['J'].width = 15.2
                ws_p.column_dimensions['K'].width = 18.9
                ws_p.column_dimensions['L'].width = 20.6

        # シート「白地図エリアマップ」生成
        ws_map = wb.create_sheet(title="白地図エリアマップ")
        ws_map.views.sheetView[0].showGridLines = True
        ws_map['A1'] = "🗺️ 大阪府 市区町村別受持選択 白地図エリアマップ（案1〜案4・最適解 比較一覧）"
        ws_map['A1'].font = Font(name='BIZ UDPゴシック', size=11, bold=True)
        
        ws_map['A3'] = "【全体トンキロ比較対比表（現状 vs 案1〜案4 vs 最適解）】"
        ws_map['A3'].font = Font(name='BIZ UDPゴシック', size=11, bold=True)
        
        map_headers = ['パターン', 'J社 トンキロ(t・km)', 'T社 トンキロ(t・km)', '合計 トンキロ(t・km)', '合計 削減量', '合計 削減率(%)']
        for col_idx, h_text in enumerate(map_headers, 1):
            c_cell = ws_map.cell(row=4, column=col_idx, value=h_text)
            c_cell.font = hdr_font
            c_cell.fill = hdr_fill
            c_cell.alignment = Alignment(horizontal='center', vertical='center')
            c_cell.border = border_thin
            
        for r_idx, c_row in enumerate(comb_summary_rows, 5):
            ws_map.cell(row=r_idx, column=1, value=c_row['パターン'])
            ws_map.cell(row=r_idx, column=2, value=c_row['J社_トンキロ'])
            ws_map.cell(row=r_idx, column=3, value=c_row['T社_トンキロ'])
            ws_map.cell(row=r_idx, column=4, value=c_row['合計_トンキロ'])
            ws_map.cell(row=r_idx, column=5, value=c_row['合計_削減量'])
            ws_map.cell(row=r_idx, column=6, value=f"{c_row['合計_削減率(%)']:+.1f}%")
            
            for c_i in range(1, 7):
                cell_item = ws_map.cell(row=r_idx, column=c_i)
                cell_item.font = data_font
                cell_item.border = border_thin

        map_cell_positions = {
            1: ('B', 13, 14),   # 案1
            2: ('E', 13, 14),   # 案2
            3: ('M', 13, 14),   # 案3
            4: ('X', 13, 14),   # 案4
            5: ('AI', 13, 14)   # 最適解
        }
        
        for p_i in [1, 2, 3, 4]:
            col_let, t_row, img_row = map_cell_positions[p_i]
            t_cell = ws_map[f"{col_let}{t_row}"]
            t_cell.value = f"■ 案{p_i} マップ"
            t_cell.font = Font(name='BIZ UDPゴシック', size=11, bold=True)
            
            m_img = pattern_map_images[p_i]
            if m_img is not None:
                b_rgb = cv2.cvtColor(m_img, cv2.COLOR_RGB2BGR)
                is_ok, buffer = cv2.imencode(".png", b_rgb)
                if is_ok:
                    img_stream = io.BytesIO(buffer)
                    xl_img = OpenPyxlImage(img_stream)
                    xl_img.width = 380
                    xl_img.height = 520
                    ws_map.add_image(xl_img, f"{col_let}{img_row}")

        col_let_opt, t_row_opt, img_row_opt = map_cell_positions[5]
        t_cell_opt = ws_map[f"{col_let_opt}{t_row_opt}"]
        t_cell_opt.value = "■ 最適解 マップ"
        t_cell_opt.font = Font(name='BIZ UDPゴシック', size=11, bold=True)
        if opt_map_image is not None:
            b_rgb_opt = cv2.cvtColor(opt_map_image, cv2.COLOR_RGB2BGR)
            is_ok_opt, buffer_opt = cv2.imencode(".png", b_rgb_opt)
            if is_ok_opt:
                xl_img_opt = OpenPyxlImage(io.BytesIO(buffer_opt))
                xl_img_opt.width = 380
                xl_img_opt.height = 520
                ws_map.add_image(xl_img_opt, f"{col_let_opt}{img_row_opt}")

        ws_map.column_dimensions['A'].width = 14.3
        ws_map.column_dimensions['B'].width = 21.5
        ws_map.column_dimensions['C'].width = 21.9
        ws_map.column_dimensions['D'].width = 21.5
        ws_map.column_dimensions['E'].width = 13.7
        ws_map.column_dimensions['F'].width = 17.0

        # シート「営業所別集計比較」の出力追加
        base_j_offices_ex = sorted(list(set(df_j['J社営業所名'].dropna().astype(str).tolist())))
        base_t_offices_ex = sorted(list(set(df_t['T社営業所名'].dropna().astype(str).tolist())))

        def get_office_metrics_dict_ex(p_df_in):
            w_dict = {}; tk_dict = {}
            for _, r_item in p_df_in.iterrows():
                bulk = str(r_item['市町村一括担当']).strip()
                bulk_off = str(r_item['市町村一括担当営業所']).strip()
                j_ind = str(r_item['J社個別']).strip()
                t_ind = str(r_item['T社個別']).strip()
                wt_j = float(r_item['J社 配達重量(kg)']) / 1000.0
                wt_t = float(r_item['T社 配達重量(kg)']) / 1000.0
                tk_j = float(r_item['J社 トンキロ(t・km)'])
                tk_t = float(r_item['T社 トンキロ(t・km)'])
                if bulk in ['J社', 'T社']:
                    key = (bulk, bulk_off)
                    w_dict[key] = w_dict.get(key, 0.0) + (wt_j + wt_t)
                    tk_dict[key] = tk_dict.get(key, 0.0) + (tk_j + tk_t)
                else:
                    if j_ind != '-':
                        key = ('J社', j_ind)
                        w_dict[key] = w_dict.get(key, 0.0) + wt_j
                        tk_dict[key] = tk_dict.get(key, 0.0) + tk_j
                    if t_ind != '-':
                        key = ('T社', t_ind)
                        w_dict[key] = w_dict.get(key, 0.0) + wt_t
                        tk_dict[key] = tk_dict.get(key, 0.0) + tk_t
            return w_dict, tk_dict

        cur_wm, cur_tkm = get_office_metrics_dict_ex(base_0_df)
        p1_wm, p1_tkm = get_office_metrics_dict_ex(st.session_state["table_data_p1"])
        p2_wm, p2_tkm = get_office_metrics_dict_ex(st.session_state["table_data_p2"])
        p3_wm, p3_tkm = get_office_metrics_dict_ex(st.session_state["table_data_p3"])
        p4_wm, p4_tkm = get_office_metrics_dict_ex(st.session_state["table_data_p4"])
        opt_wm, opt_tkm = get_office_metrics_dict_ex(best_opt_res['df'])

        def build_comp_ex_df(comp_name, base_offs, is_weight=True):
            rows = []
            m_cur, m_p1, m_p2, m_p3, m_p4, m_opt = (cur_wm, p1_wm, p2_wm, p3_wm, p4_wm, opt_wm) if is_weight else (cur_tkm, p1_tkm, p2_tkm, p3_tkm, p4_tkm, opt_tkm)
            suffix = "_配達重量_t" if is_weight else "_トンキロ"
            
            for off in base_offs:
                rows.append({
                    '営業所名': off,
                    f'現状{suffix}': round(m_cur.get((comp_name, off), 0.0), 2),
                    f'案1{suffix}': round(m_p1.get((comp_name, off), 0.0), 2),
                    f'案2{suffix}': round(m_p2.get((comp_name, off), 0.0), 2),
                    f'案3{suffix}': round(m_p3.get((comp_name, off), 0.0), 2),
                    f'案4{suffix}': round(m_p4.get((comp_name, off), 0.0), 2),
                    f'最適解{suffix}': round(m_opt.get((comp_name, off), 0.0), 2),
                })
            
            p1_n = round(m_p1.get((comp_name, f"{comp_name}案1"), 0.0), 2)
            p2_n = round(m_p2.get((comp_name, f"{comp_name}案2"), 0.0), 2)
            p3_n = round(m_p3.get((comp_name, f"{comp_name}案3"), 0.0), 2)
            p4_n = round(m_p4.get((comp_name, f"{comp_name}案4"), 0.0), 2)
            opt_n = round(m_opt.get((comp_name, f"{comp_name}最適"), 0.0), 2)
            
            rows.append({
                '営業所名': f"{comp_name}新規営業所",
                f'現状{suffix}': 0.0,
                f'案1{suffix}': p1_n, f'案2{suffix}': p2_n, f'案3{suffix}': p3_n, f'案4{suffix}': p4_n, f'最適解{suffix}': opt_n,
            })
            
            df_sub = pd.DataFrame(rows)
            tot_row = {'営業所名': '合計'}
            for col_c in df_sub.columns:
                if col_c != '営業所名':
                    tot_row[col_c] = round(df_sub[col_c].sum(), 2)
                    
            return pd.concat([df_sub, pd.DataFrame([tot_row])], ignore_index=True)

        ws_off = wb.create_sheet(title="営業所別集計比較")
        ws_off.views.sheetView[0].showGridLines = True
        
        ws_off['A1'] = "【1. 各営業所の「現状」vs「案1〜4」vs「最適解」合計配達重量（t）比較】"
        ws_off['A1'].font = Font(name='BIZ UDPゴシック', size=11, bold=True)
        
        df_j_w_ex = build_comp_ex_df('J社', base_j_offices_ex, is_weight=True)
        df_t_w_ex = build_comp_ex_df('T社', base_t_offices_ex, is_weight=True)
        
        ws_off['A3'] = "🔴 J社 営業所別 配達重量（t）比較表"
        ws_off['A3'].font = Font(name='BIZ UDPゴシック', size=11, bold=True)
        for col_i, col_name in enumerate(df_j_w_ex.columns, 1):
            cell_h = ws_off.cell(row=4, column=col_i, value=col_name)
            cell_h.font = hdr_font; cell_h.fill = hdr_fill; cell_h.border = border_thin
        for r_i, r_val in df_j_w_ex.iterrows():
            for c_i, col_name in enumerate(df_j_w_ex.columns, 1):
                cell_d = ws_off.cell(row=r_i+5, column=c_i, value=r_val[col_name])
                cell_d.font = data_font; cell_d.border = border_thin

        row_start_t = len(df_j_w_ex) + 7
        ws_off.cell(row=row_start_t, column=1, value="🔵 T社 営業所別 配達重量（t）比較表").font = Font(name='BIZ UDPゴシック', size=11, bold=True)
        for col_i, col_name in enumerate(df_t_w_ex.columns, 1):
            cell_h = ws_off.cell(row=row_start_t+1, column=col_i, value=col_name)
            cell_h.font = hdr_font; cell_h.fill = hdr_fill; cell_h.border = border_thin
        for r_i, r_val in df_t_w_ex.iterrows():
            for c_i, col_name in enumerate(df_t_w_ex.columns, 1):
                cell_d = ws_off.cell(row=row_start_t+2+r_i, column=c_i, value=r_val[col_name])
                cell_d.font = data_font; cell_d.border = border_thin

        row_start_tk = row_start_t + len(df_t_w_ex) + 4
        ws_off.cell(row=row_start_tk, column=1, value="【2. 各営業所の「現状」vs「案1〜4」vs「最適解」合計トンキロ（t・km）比較】").font = Font(name='BIZ UDPゴシック', size=11, bold=True)
        
        df_j_tk_ex = build_comp_ex_df('J社', base_j_offices_ex, is_weight=False)
        df_t_tk_ex = build_comp_ex_df('T社', base_t_offices_ex, is_weight=False)

        ws_off.cell(row=row_start_tk+2, column=1, value="🔴 J社 営業所別 トンキロ（t・km）比較表").font = Font(name='BIZ UDPゴシック', size=11, bold=True)
        for col_i, col_name in enumerate(df_j_tk_ex.columns, 1):
            cell_h = ws_off.cell(row=row_start_tk+3, column=col_i, value=col_name)
            cell_h.font = hdr_font; cell_h.fill = hdr_fill; cell_h.border = border_thin
        for r_i, r_val in df_j_tk_ex.iterrows():
            for c_i, col_name in enumerate(df_j_tk_ex.columns, 1):
                cell_d = ws_off.cell(row=row_start_tk+4+r_i, column=c_i, value=r_val[col_name])
                cell_d.font = data_font; cell_d.border = border_thin

        row_start_t2 = row_start_tk + len(df_j_tk_ex) + 6
        ws_off.cell(row=row_start_t2, column=1, value="🔵 T社 営業所別 トンキロ（t・km）比較表").font = Font(name='BIZ UDPゴシック', size=11, bold=True)
        for col_i, col_name in enumerate(df_t_tk_ex.columns, 1):
            cell_h = ws_off.cell(row=row_start_t2+1, column=col_i, value=col_name)
            cell_h.font = hdr_font; cell_h.fill = hdr_fill; cell_h.border = border_thin
        for r_i, r_val in df_t_tk_ex.iterrows():
            for c_i, col_name in enumerate(df_t_tk_ex.columns, 1):
                cell_d = ws_off.cell(row=row_start_t2+2+r_i, column=c_i, value=r_val[col_name])
                cell_d.font = data_font; cell_d.border = border_thin

        for c_i in range(1, 8):
            col_let_o = openpyxl.utils.get_column_letter(c_i)
            ws_off.column_dimensions[col_let_o].width = 18.0

    st.download_button(
        label="📥 全体シミュレーション＆一括データ（現状・案1〜4・最適解・白地図・営業所別集計併記）Excel出力",
        data=excel_full_bytes.getvalue(),
        file_name="大阪府配達エリア最適化シミュレーション_全体比較結果.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        type="primary"
    )

    st.markdown("---")
    # J社＋T社「全体合計」トンキロ比較メトリック（最適解含む6パターン）
    st.markdown("### 🚚 J社 ＋ T社 全体合計：現状 vs 4案 vs 最適解のトンキロ比較")
    cols_comb = st.columns(6)
    for idx, col in enumerate(cols_comb):
        row_c = comb_summary_df.iloc[idx]
        diff_v = row_c['合計_削減量']
        pct_v = row_c['合計_削減率(%)']
        
        if idx == 0:
            delta_label = None
        else:
            if diff_v >= 0:
                delta_label = f"削減: {diff_v:,.1f} ({pct_v:+.1f}%)"
            else:
                inc_val = abs(diff_v)
                inc_pct = abs(pct_v)
                delta_label = f"増加: +{inc_val:,.1f} (+{inc_pct:.1f}%)"

        with col:
            st.markdown(f"#### {row_c['パターン']}")
            st.metric("合計トンキロ", "{:,.1f} ton-km".format(row_c['合計_トンキロ']), delta=delta_label)
            st.caption(f"J社: {row_c['J社_トンキロ']:,.1f} ton-km / T社: {row_c['T社_トンキロ']:,.1f} ton-km")

    st.markdown("---")
    st.write("### 📊 全パターン比較対比表（会社別・最適解含む）")
    st.dataframe(
        summary_df.style.format({
            '配達重量_t': '{:,.2f}', 'トンキロ': '{:,.2f}',
            'トンキロ削減量': '{:,.2f}', '削減率(%)': '{:+.1f}%'
        }),
        use_container_width=True, hide_index=True
    )

# タブ2: 営業所別 集計比較（配達重量 ＆ トンキロ に「最適解」を追加・統合）
with tabs[1]:
    st.subheader("🏬 各営業所の「現状」vs「案1〜案4」vs「最適解」集計比較")
    st.caption("J社・T社それぞれの営業所について、「現状」「案1〜案4」「最適解」での合計配達重量（t）および合計トンキロ（t・km）を一括比較します。")
    
    base_j_offices = sorted(list(set(df_j['J社営業所名'].dropna().astype(str).tolist())))
    base_t_offices = sorted(list(set(df_t['T社営業所名'].dropna().astype(str).tolist())))

    def get_office_metrics_dict(p_df_in):
        w_dict = {}
        tk_dict = {}
        for _, r_item in p_df_in.iterrows():
            bulk = str(r_item['市町村一括担当']).strip()
            bulk_off = str(r_item['市町村一括担当営業所']).strip()
            j_ind = str(r_item['J社個別']).strip()
            t_ind = str(r_item['T社個別']).strip()
            
            wt_j = float(r_item['J社 配達重量(kg)']) / 1000.0
            wt_t = float(r_item['T社 配達重量(kg)']) / 1000.0
            tk_j = float(r_item['J社 トンキロ(t・km)'])
            tk_t = float(r_item['T社 トンキロ(t・km)'])
            
            if bulk in ['J社', 'T社']:
                key = (bulk, bulk_off)
                w_dict[key] = w_dict.get(key, 0.0) + (wt_j + wt_t)
                tk_dict[key] = tk_dict.get(key, 0.0) + (tk_j + tk_t)
            else:
                if j_ind != '-':
                    key = ('J社', j_ind)
                    w_dict[key] = w_dict.get(key, 0.0) + wt_j
                    tk_dict[key] = tk_dict.get(key, 0.0) + tk_j
                if t_ind != '-':
                    key = ('T社', t_ind)
                    w_dict[key] = w_dict.get(key, 0.0) + wt_t
                    tk_dict[key] = tk_dict.get(key, 0.0) + tk_t
        return w_dict, tk_dict

    cur_w_map, cur_tk_map = get_office_metrics_dict(build_initial_scenario_df(EXISTING_OFFICES_INFO))
    p1_w_map, p1_tk_map = get_office_metrics_dict(st.session_state["table_data_p1"])
    p2_w_map, p2_tk_map = get_office_metrics_dict(st.session_state["table_data_p2"])
    p3_w_map, p3_tk_map = get_office_metrics_dict(st.session_state["table_data_p3"])
    p4_w_map, p4_tk_map = get_office_metrics_dict(st.session_state["table_data_p4"])
    opt_w_map, opt_tk_map = get_office_metrics_dict(best_opt_res['df'])

    # 1. 合計配達重量（t）比較テーブルの構築
    def build_company_weight_compare_df(comp_name, base_offices):
        rows = []
        for off in base_offices:
            rows.append({
                '営業所名': off,
                '現状_配達重量_t': round(cur_w_map.get((comp_name, off), 0.0), 2),
                '案1_配達重量_t': round(p1_w_map.get((comp_name, off), 0.0), 2),
                '案2_配達重量_t': round(p2_w_map.get((comp_name, off), 0.0), 2),
                '案3_配達重量_t': round(p3_w_map.get((comp_name, off), 0.0), 2),
                '案4_配達重量_t': round(p4_w_map.get((comp_name, off), 0.0), 2),
                '最適解_配達重量_t': round(opt_w_map.get((comp_name, off), 0.0), 2),
            })
            
        new_off_row_name = f"{comp_name}新規営業所"
        p1_new_val = round(p1_w_map.get((comp_name, f"{comp_name}案1"), 0.0), 2)
        p2_new_val = round(p2_w_map.get((comp_name, f"{comp_name}案2"), 0.0), 2)
        p3_new_val = round(p3_w_map.get((comp_name, f"{comp_name}案3"), 0.0), 2)
        p4_new_val = round(p4_w_map.get((comp_name, f"{comp_name}案4"), 0.0), 2)
        opt_new_val = round(opt_w_map.get((comp_name, f"{comp_name}最適"), 0.0), 2)
        
        rows.append({
            '営業所名': new_off_row_name,
            '現状_配達重量_t': 0.0,
            '案1_配達重量_t': p1_new_val,
            '案2_配達重量_t': p2_new_val,
            '案3_配達重量_t': p3_new_val,
            '案4_配達重量_t': p4_new_val,
            '最適解_配達重量_t': opt_new_val,
        })
        
        df_sub = pd.DataFrame(rows)
        tot_row = {
            '営業所名': '合計',
            '現状_配達重量_t': round(df_sub['現状_配達重量_t'].sum(), 2),
            '案1_配達重量_t': round(df_sub['案1_配達重量_t'].sum(), 2),
            '案2_配達重量_t': round(df_sub['案2_配達重量_t'].sum(), 2),
            '案3_配達重量_t': round(df_sub['案3_配達重量_t'].sum(), 2),
            '案4_配達重量_t': round(df_sub['案4_配達重量_t'].sum(), 2),
            '最適解_配達重量_t': round(df_sub['最適解_配達重量_t'].sum(), 2),
        }
        return pd.concat([df_sub, pd.DataFrame([tot_row])], ignore_index=True)

    # 2. 合計トンキロ（t・km）比較テーブルの構築
    def build_company_tonkm_compare_df(comp_name, base_offices):
        rows = []
        for off in base_offices:
            rows.append({
                '営業所名': off,
                '現状_トンキロ': round(cur_tk_map.get((comp_name, off), 0.0), 2),
                '案1_トンキロ': round(p1_tk_map.get((comp_name, off), 0.0), 2),
                '案2_トンキロ': round(p2_tk_map.get((comp_name, off), 0.0), 2),
                '案3_トンキロ': round(p3_tk_map.get((comp_name, off), 0.0), 2),
                '案4_トンキロ': round(p4_tk_map.get((comp_name, off), 0.0), 2),
                '最適解_トンキロ': round(opt_tk_map.get((comp_name, off), 0.0), 2),
            })
            
        new_off_row_name = f"{comp_name}新規営業所"
        p1_new_val = round(p1_tk_map.get((comp_name, f"{comp_name}案1"), 0.0), 2)
        p2_new_val = round(p2_tk_map.get((comp_name, f"{comp_name}案2"), 0.0), 2)
        p3_new_val = round(p3_tk_map.get((comp_name, f"{comp_name}案3"), 0.0), 2)
        p4_new_val = round(p4_tk_map.get((comp_name, f"{comp_name}案4"), 0.0), 2)
        opt_new_val = round(opt_tk_map.get((comp_name, f"{comp_name}最適"), 0.0), 2)
        
        rows.append({
            '営業所名': new_off_row_name,
            '現状_トンキロ': 0.0,
            '案1_トンキロ': p1_new_val,
            '案2_トンキロ': p2_new_val,
            '案3_トンキロ': p3_new_val,
            '案4_トンキロ': p4_new_val,
            '最適解_トンキロ': opt_new_val,
        })
        
        df_sub = pd.DataFrame(rows)
        tot_row = {
            '営業所名': '合計',
            '現状_トンキロ': round(df_sub['現状_トンキロ'].sum(), 2),
            '案1_トンキロ': round(df_sub['案1_トンキロ'].sum(), 2),
            '案2_トンキロ': round(df_sub['案2_トンキロ'].sum(), 2),
            '案3_トンキロ': round(df_sub['案3_トンキロ'].sum(), 2),
            '案4_トンキロ': round(df_sub['案4_トンキロ'].sum(), 2),
            '最適解_トンキロ': round(df_sub['最適解_トンキロ'].sum(), 2),
        }
        return pd.concat([df_sub, pd.DataFrame([tot_row])], ignore_index=True)

    # 1) 配達重量（t）比較表示
    st.markdown("### 📦 1. 各営業所の「現状」vs「案1〜案4」vs「最適解」合計配達重量（t）比較")
    col_w1, col_w2 = st.columns(2)
    with col_w1:
        st.markdown("#### 🔴 J社 営業所別 配達重量（t）比較表")
        st.dataframe(
            build_company_weight_compare_df('J社', base_j_offices).style.format({
                '現状_配達重量_t': '{:,.2f}', '案1_配達重量_t': '{:,.2f}',
                '案2_配達重量_t': '{:,.2f}', '案3_配達重量_t': '{:,.2f}',
                '案4_配達重量_t': '{:,.2f}', '最適解_配達重量_t': '{:,.2f}'
            }),
            use_container_width=True, hide_index=True
        )

    with col_w2:
        st.markdown("#### 🔵 T社 営業所別 配達重量（t）比較表")
        st.dataframe(
            build_company_weight_compare_df('T社', base_t_offices).style.format({
                '現状_配達重量_t': '{:,.2f}', '案1_配達重量_t': '{:,.2f}',
                '案2_配達重量_t': '{:,.2f}', '案3_配達重量_t': '{:,.2f}',
                '案4_配達重量_t': '{:,.2f}', '最適解_配達重量_t': '{:,.2f}'
            }),
            use_container_width=True, hide_index=True
        )

    st.markdown("---")

    # 2) トンキロ（t・km）比較表示
    st.markdown("### 🚚 2. 各営業所の「現状」vs「案1〜案4」vs「最適解」合計トンキロ（t・km）比較")
    col_tk1, col_tk2 = st.columns(2)
    with col_tk1:
        st.markdown("#### 🔴 J社 営業所別 トンキロ（t・km）比較表")
        st.dataframe(
            build_company_tonkm_compare_df('J社', base_j_offices).style.format({
                '現状_トンキロ': '{:,.2f}', '案1_トンキロ': '{:,.2f}',
                '案2_トンキロ': '{:,.2f}', '案3_トンキロ': '{:,.2f}',
                '案4_トンキロ': '{:,.2f}', '最適解_トンキロ': '{:,.2f}'
            }),
            use_container_width=True, hide_index=True
        )

    with col_tk2:
        st.markdown("#### 🔵 T社 営業所別 トンキロ（t・km）比較表")
        st.dataframe(
            build_company_tonkm_compare_df('T社', base_t_offices).style.format({
                '現状_トンキロ': '{:,.2f}', '案1_トンキロ': '{:,.2f}',
                '案2_トンキロ': '{:,.2f}', '案3_トンキロ': '{:,.2f}',
                '案4_トンキロ': '{:,.2f}', '最適解_トンキロ': '{:,.2f}'
            }),
            use_container_width=True, hide_index=True
        )